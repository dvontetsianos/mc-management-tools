#one-off tool: imports Eirini's Housekeeping catalogue into the Housekeeping department
#
#usage (from the project folder, the Excel file is on the Desktop next to it):
#   python import_hk_catalogue.py ..\HK_database_List_Draft_XLS.xlsx            shows what would happen, changes NOTHING
#   python import_hk_catalogue.py ..\HK_database_List_Draft_XLS.xlsx --apply    backs up the database first, then imports
#
#what it does, per row of the file:
#   category       her section (BED LINENS...), created in Housekeeping if missing
#   subcategory    her Sub-Category column, created under that category if missing
#   name           her Specification / Description column
#   code, material, size, color, supplier description   into the catalogue fields
#   hotels         all hotels, opening quantity 0
#
#rows that are skipped:
#   empty rows (a code with nothing else, e.g. HK-INR-14)
#   codes that already exist in the database (so running it twice never doubles anything)
#
#needs migrate_item_names_with_code.py first, because many items share a name
#everything is done in ONE save: if anything goes wrong, nothing changes

import argparse
import re
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DB_FILE = BASE_DIR / "assets.db"

DEPARTMENT_NAME = "Housekeeping"

#her section titles -> our category names
SECTIONS = {
    "BED LINENS": "Bed Linens",
    "BATH LINENS": "Bath Linens",
    "SPA & GYM LINENS": "Spa & Gym Linens",
    "IN ROOM (NON CONSUMABLES": "In Room (Non Consumables)",
    "HOUSEKEEPING & LAUNDRY EQUIPMENT": "Housekeeping & Laundry Equipment",
}


#---------------------------------------------------------------------------
#one clean line of text: no double spaces, no line breaks, empty -> None
def clean(value):

    if value is None:
        return None

    text = " ".join(str(value).replace("\xa0", " ").split())

    return text or None


#sizes: Greek Χ and small x between numbers become X (80Χ150 -> 80X150), so searching always works
#Excel turned some sizes into dates (4-6 -> 4 June), those are turned back
def clean_size(value):

    if isinstance(value, datetime):
        return f"{value.day}-{value.month}"

    text = clean(value)

    if text is None:
        return None

    return re.sub(r"(?<=\d)\s*[ΧχXx]\s*(?=\d)", "X", text)


#---------------------------------------------------------------------------
#reads the file into a list of rows, and a list of problems
def read_file(path):

    sheet = load_workbook(path, data_only=True).active

    rows = []
    skipped = []
    problems = []
    category = None

    for line_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = list(values) + [None] * 9
        first = clean(values[0])

        if first is None:
            continue

        #a section title, e.g. BED LINENS: the rows below it belong to that category
        if not first.upper().startswith("HK-"):
            if first.upper() in SECTIONS:
                category = SECTIONS[first.upper()]
            continue

        code = first.upper()
        name = clean(values[3]) or clean(values[2])

        if name is None:
            skipped.append(f"{code} (empty row)")
            continue

        if category is None:
            problems.append(f"line {line_number}: {code} comes before any section title")
            continue

        rows.append({
            "line": line_number,
            "code": code,
            "category": category,
            "subcategory": clean(values[1]),
            "name": name,
            "material": clean(values[4]),
            "size": clean_size(values[5]),
            "color": clean(values[6]),
            "supplier_description": clean(values[7]),
        })

    seen = {}
    for row in rows:
        if row["code"] in seen:
            problems.append(f"line {row['line']}: code {row['code']} is also on line {seen[row['code']]}")
        seen[row["code"]] = row["line"]

    return rows, skipped, problems


#---------------------------------------------------------------------------
#finds a category of the department, or creates it
def get_category_id(db, department_id, name, created):

    row = db.execute(
        "SELECT id FROM categories WHERE department_id = ? AND name = ?", (department_id, name)
    ).fetchone()

    if row:
        return row["id"]

    created.append(f"category {name}")
    return db.execute(
        "INSERT INTO categories (name, department_id) VALUES (?, ?)", (name, department_id)
    ).lastrowid


#finds a subcategory of the category, or creates it
def get_subcategory_id(db, category_id, name, created):

    if name is None:
        return None

    row = db.execute(
        "SELECT id FROM subcategories WHERE category_id = ? AND name = ?", (category_id, name)
    ).fetchone()

    if row:
        return row["id"]

    created.append(f"subcategory {name}")
    return db.execute(
        "INSERT INTO subcategories (name, category_id) VALUES (?, ?)", (name, category_id)
    ).lastrowid


#---------------------------------------------------------------------------
def main():

    parser = argparse.ArgumentParser()
    parser.add_argument("excel_file", help="Eirini's HK catalogue (.xlsx)")
    parser.add_argument("--apply", action="store_true", help="really import (without it nothing changes)")
    args = parser.parse_args()

    if not DB_FILE.exists():
        raise SystemExit(f"No database found at {DB_FILE} - run this from the project folder.")

    rows, skipped, problems = read_file(args.excel_file)

    db = sqlite3.connect(DB_FILE)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys = ON")

    department = db.execute("SELECT id FROM departments WHERE name = ?", (DEPARTMENT_NAME,)).fetchone()
    if department is None:
        problems.append(f"department {DEPARTMENT_NAME} doesn't exist")

    name_rule = db.execute(
        "SELECT sql FROM sqlite_master WHERE type = 'index' AND name = 'unique_item_name_per_department'"
    ).fetchone()
    if name_rule and "WHERE" not in name_rule["sql"].upper():
        problems.append("run migrate_item_names_with_code.py first (many items share a name)")

    hotel_ids = [row["id"] for row in db.execute("SELECT id FROM hotels ORDER BY id")]
    if not hotel_ids:
        problems.append("there are no hotels in the database")

    if problems:
        print("NOTHING WAS CHANGED, fix these first:")
        for problem in problems:
            print("  - " + problem)
        sys.exit(1)

    #the same backup update.bat makes, and stop if it fails
    if args.apply:
        print("Backing up the database...")
        backup = subprocess.run([sys.executable, str(BASE_DIR / "backup_db.py"), "--before-update"], cwd=BASE_DIR)
        if backup.returncode != 0:
            print("BACKUP FAILED - nothing was changed.")
            sys.exit(1)

    created = []
    imported = 0
    already = []
    now = str(datetime.utcnow())

    try:
        for row in rows:

            if db.execute("SELECT id FROM items WHERE code = ?", (row["code"],)).fetchone():
                already.append(row["code"])
                continue

            category_id = get_category_id(db, department["id"], row["category"], created)
            subcategory_id = get_subcategory_id(db, category_id, row["subcategory"], created)

            item_id = db.execute(
                """INSERT INTO items (name, category_id, subcategory_id, department_id, opening_quantity,
                                      code, material, size, color, supplier_description)
                   VALUES (?, ?, ?, ?, 0, ?, ?, ?, ?, ?)""",
                (
                    row["name"], category_id, subcategory_id, department["id"],
                    row["code"], row["material"], row["size"], row["color"], row["supplier_description"],
                ),
            ).lastrowid

            for hotel_id in hotel_ids:
                db.execute(
                    "INSERT INTO item_hotels (item_id, hotel_id, opening_quantity) VALUES (?, ?, 0)",
                    (item_id, hotel_id),
                )

            db.execute(
                """INSERT INTO actions_logs (username, action, entity_type, entity_name, entity_id, details, timestamp)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                ("import script", "created", "Item", row["name"], item_id,
                 f"Imported {row['code']} from the Housekeeping catalogue", now),
            )

            imported += 1
            print(f"  {row['code']:<11} {row['category']:<34} {row['name']}"
                  f"{'  [' + row['size'] + ']' if row['size'] else ''}")

        for line in created:
            print("  new " + line)

        if skipped:
            print(f"Skipped (empty in the file): {', '.join(skipped)}")
        if already:
            print(f"Skipped (code already in the database): {len(already)}")

        if args.apply:
            db.commit()
            print(f"\nDONE: {imported} item(s) imported into {DEPARTMENT_NAME}, linked to {len(hotel_ids)} hotel(s).")
        else:
            db.rollback()
            print(f"\nTEST RUN ONLY: nothing was changed. Run again with --apply to import these {imported} item(s).")

    except Exception as error:
        db.rollback()
        print(f"\nSOMETHING WENT WRONG, nothing was changed: {error}")
        sys.exit(1)

    finally:
        db.close()


if __name__ == "__main__":
    main()