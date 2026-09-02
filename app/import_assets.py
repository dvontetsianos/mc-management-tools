"""
Import assets from the "Ready For Import" sheet of asset_import_candidates.xlsx
into the FB Assets database.

Run from the project root (the FB_Assets folder, same place assets.db lives):

    python -m app.import_assets

ALWAYS run it once with DRY_RUN = True first and read the summary before
changing DRY_RUN to False and running it again for real.
"""


from openpyxl import load_workbook
from .database import SessionLocal
from . import models


EXCEL_PATH = r"C:\Users\d.vontetsianos\Desktop\asset_import_candidates.xlsx"
SHEET_NAME = "Ready For Import"


DRY_RUN = True #change to false once you verify the summary


def read_rows():
    wb = load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    rows = []
    seen = {}

    for row in ws.iter_rows(min_row=5, values_only=True):
        name, category, location, qty, broken, notes = row


        if not name or not category:
            continue #section headers/instructions/blank rows


        name = str(name).strip()
        category = str(category).strip()
        location = str(location).strip() if location else None
        qty = int(qty) if isinstance(qty, (int, float)) else 0
        broken = int(broken) if isinstance(broken, (int, float)) else 0


        if not location:
            print(f"Skipped, no location: {name}")
            continue

        key = (name, location)
        if key in seen:
            prev_qty, prev_broken = seen[key]
            if (prev_qty, prev_broken) != (qty, broken):
                print(
                    f"Warning: duplicate row for {key} with different values "
                    f"({prev_qty}/{prev_broken} vs {qty}/{broken}) - keeping the first, check manually."
                )

            continue #keeps first only

        seen[key] = (qty, broken)
        rows.append({
            "name": name,
            "category": category,
            "location": location,
            "total_quantity": qty,
            "broken_quantity": broken,
        })

    return rows



def get_or_create_category(db, name, cache):
    key = name.lower()
    if key in cache:
        return cache[key], False


    existing = db.query(models.Category).filter(models.Category.name.ilike(name)).first()
    if existing:
        cache[key] = existing
        return existing, False

    new_cat = models.Category(name=name)
    db.add(new_cat)
    db.flush()
    cache[key] = new_cat
    return new_cat, True


def get_or_create_location(db, name, cache):
    key = name.lower()
    if key in cache:
        return cache[key], False


    existing = db.query(models.Location).filter(models.Location.name.ilike(name)).first()
    if existing:
        cache[key] = existing
        return existing, False


    new_loc = models.Location(name=name)
    db.add(new_loc)
    db.flush()
    cache[key] = new_loc
    return new_loc, True


def main():
    rows = read_rows()
    print(f"Parsed {len(rows)} clean rows from the spreadsheet. \n")

    db = SessionLocal()

    category_cache = {}
    location_cache = {}

    categories_created = 0
    locations_created = 0
    assets_created = 0
    assets_skipped = 0


    for row in rows:
        existing_asset = (
            db.query(models.Asset)
            .join(models.Location)
            .filter(models.Asset.name == row["name"])
            .filter(models.Location.name == row["location"])
            .first()
        )
        if existing_asset:
            assets_skipped += 1
            continue


        category, cat_created = get_or_create_category(db, row["category"], category_cache)
        if cat_created:
            categories_created += 1


        location, loc_created = get_or_create_location(db, row["location"], location_cache)
        if loc_created:
            locations_created += 1


        asset = models.Asset(
            name=row["name"],
            category_id=category.id,
            location_id=location.id,
            total_quantity=row["total_quantity"],
            broken_quantity=row["broken_quantity"],
        )
        db.add(asset)
        assets_created += 1


    print("---- SUMMARY ----")
    print(f"Categories to create: {categories_created}")
    print(f"Locations to create:  {locations_created}")
    print(f"Assets to create:     {assets_created}")
    print(f"Assets skipped        {assets_skipped}")



    if DRY_RUN:
        db.rollback()
        print("\nDry RUN - nothing was written to the database.")
    else:
        db.commit()
        print("\nDONE - changes commited to the database.")

    db.close()


if __name__ == "__main__":
    main()